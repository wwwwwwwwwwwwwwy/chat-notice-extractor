"""test_pipeline.py - chat-notice-extractor 端到端测试

运行：
    python -m unittest discover -s tests -v
或：
    python tests/test_pipeline.py
"""

import json
import os
import sys
import tempfile
import unittest
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import build_table  # noqa: E402
import demo_data  # noqa: E402
import notice_rules  # noqa: E402
import scan_notices  # noqa: E402
import xlsx_min  # noqa: E402
from notice_rules import (  # noqa: E402
    TZ,
    build_candidates,
    infer_chat_type,
    is_noise,
    is_notice_like,
    score_message,
    ts_to_str,
)


def make_msg(local_id, content, sender="them", msg_type="text", ts=1772419680, name="班长"):
    return {
        "local_id": local_id,
        "sender": sender,
        "sender_name": name,
        "content": content,
        "timestamp": ts,
        "type": msg_type,
    }


class TestTimestamp(unittest.TestCase):
    def test_seconds(self):
        dt = datetime(2026, 3, 2, 9, 28, 0, tzinfo=TZ)
        self.assertEqual(ts_to_str(int(dt.timestamp())), "2026-03-02 09:28:00")

    def test_milliseconds_are_downgraded(self):
        dt = datetime(2026, 3, 2, 9, 28, 0, tzinfo=TZ)
        seconds = int(dt.timestamp())
        self.assertEqual(ts_to_str(seconds * 1000), "2026-03-02 09:28:00")
        self.assertEqual(ts_to_str(seconds * 1000 * 1000), "2026-03-02 09:28:00")

    def test_invalid_returns_empty(self):
        for bad in (None, "", "abc", {}, []):
            self.assertEqual(ts_to_str(bad), "")

    def test_timezone_is_utc_plus_8(self):
        # 1970-01-01 00:00:00 UTC == 1970-01-01 08:00:00 +08:00
        self.assertEqual(ts_to_str(0), "1970-01-01 08:00:00")


class TestNoise(unittest.TestCase):
    def test_short_replies_are_noise(self):
        for text in ("嗯", "哦", "好的", "收到", "哈哈", "在", "?"):
            self.assertTrue(is_noise(text, "text"), text)

    def test_media_types_are_noise(self):
        for msg_type in ("image", "voice", "video", "emoji", "revoke", "call"):
            self.assertTrue(is_noise("随便什么内容", msg_type), msg_type)

    def test_real_content_is_not_noise(self):
        self.assertFalse(is_noise("通知一下：明天下午三点开会", "text"))


class TestScoring(unittest.TestCase):
    def test_notice_scores_high(self):
        msg = make_msg(1, "通知一下：本周五下午3点在B301开班会，请全体成员务必参加。")
        self.assertGreaterEqual(score_message(msg)["score"], 6)

    def test_noise_scores_zero(self):
        self.assertEqual(score_message(make_msg(1, "嗯"))["score"], 0)

    def test_greeting_gate_fails(self):
        """只有泛称呼 + 时段，不应判定为通知。"""
        msg = make_msg(1, "同学们早上好")
        result = score_message(msg)
        self.assertFalse(is_notice_like(result["score"], result["hits"], 3))

    def test_time_only_gate_fails(self):
        """只有「时段」不足以判定为通知。"""
        msg = make_msg(1, "晚上一起吃饭吗")
        result = score_message(msg)
        self.assertFalse(is_notice_like(result["score"], result["hits"], 3))

    def test_specific_time_gate_passes(self):
        """具体时间（X点）可独立支撑判定。"""
        msg = make_msg(1, "明天下午三点老地方见吧，我把合同带过去。")
        result = score_message(msg)
        self.assertTrue(is_notice_like(result["score"], result["hits"], 3))

    def test_below_threshold_never_passes(self):
        msg = make_msg(1, "通知明天开会")  # 有关键词但分数低
        result = score_message(msg)
        self.assertFalse(is_notice_like(result["score"], result["hits"], 99))


class TestChatType(unittest.TestCase):
    def test_declared_wins(self):
        payload = {"chat_type": "group"}
        self.assertEqual(infer_chat_type(payload, []), "group")
        self.assertEqual(infer_chat_type({"chat_type": "private"}, []), "private")

    def test_group_notice_marker(self):
        self.assertEqual(infer_chat_type({}, [make_msg(1, "群公告：明天放假")]), "group")

    def test_at_all_marker(self):
        self.assertEqual(infer_chat_type({}, [make_msg(1, "@全体成员 注意")]), "group")

    def test_unknown_when_no_signal(self):
        self.assertEqual(infer_chat_type({}, [make_msg(1, "在吗")]), "unknown")


class TestCandidates(unittest.TestCase):
    def test_anchor_ids_subset_of_window(self):
        payload = {"contact_display": "测试群", "chat_type": "group",
                   "messages": [make_msg(i, "通知：明天下午3点开会，请全体成员参加。") for i in range(1, 6)]}
        result = build_candidates(payload)
        for candidate in result["candidates"]:
            window_ids = {m["local_id"] for m in candidate["window"]}
            for anchor_id in candidate["anchor_local_ids"]:
                self.assertIn(anchor_id, window_ids)

    def test_adjacent_notices_merge(self):
        """相邻两条通知消息应合并为一个候选。"""
        messages = [
            make_msg(1, "通知：明天下午3点开会，请全体成员参加。"),
            make_msg(2, "补充：地点改到B301，务必准时。"),
        ]
        result = build_candidates({"contact_display": "x", "messages": messages})
        self.assertEqual(len(result["candidates"]), 1)

    def test_distant_notices_stay_separate(self):
        """相隔较远的通知不应被并成一条。"""
        messages = [
            make_msg(1, "通知：明天下午3点开会，请全体成员参加。"),
            make_msg(2, "哈哈"),
            make_msg(3, "好的"),
            make_msg(4, "收到"),
            make_msg(5, "另外接龙报名春游，下周六去青龙峡，报名截止到3月10日晚上8点。"),
        ]
        result = build_candidates({"contact_display": "x", "messages": messages})
        self.assertEqual(len(result["candidates"]), 2)

    def test_context_expands_window(self):
        messages = [make_msg(1, "闲聊")] + [
            make_msg(i, "通知：明天下午3点开会，请全体成员参加。") for i in range(2, 5)
        ] + [make_msg(5, "闲聊结尾")]
        result = build_candidates({"contact_display": "x", "messages": messages}, context=2)
        window_ids = {m["local_id"] for m in result["candidates"][0]["window"]}
        self.assertIn(1, window_ids)
        self.assertIn(5, window_ids)

    def test_empty_messages(self):
        result = build_candidates({"contact_display": "x", "messages": []})
        self.assertEqual(result["candidates"], [])

    def test_no_anchor_returns_empty(self):
        messages = [make_msg(1, "嗯"), make_msg(2, "好的"), make_msg(3, "哈哈")]
        result = build_candidates({"contact_display": "x", "messages": messages})
        self.assertEqual(result["candidates"], [])

    def test_candidate_ids_are_unique_after_scan(self):
        payload = {"contact_display": "x", "messages": [
            make_msg(1, "通知：明天下午3点开会，请全体成员参加。"),
            make_msg(5, "接龙报名春游，下周六出发，报名截止到3月10日。"),
        ]}
        result = build_candidates(payload)
        for index, candidate in enumerate(result["candidates"], start=1):
            candidate["candidate_id"] = "c%04d" % index
        ids = [c["candidate_id"] for c in result["candidates"]]
        self.assertEqual(len(ids), len(set(ids)))


class TestBestWindowMessage(unittest.TestCase):
    """回归测试：窗口重叠时，必须选锚点消息，而不是邻近的另一条通知。"""

    def test_prefers_anchor_over_neighbor(self):
        neighbour = make_msg(1, "通知一下：本周五下午3点在B301开班会，请全体成员务必参加。",
                             ts=1772419680)
        anchor = make_msg(2, "哦对了，周五之前得把报价单发我邮箱，客户催得比较紧。",
                          ts=1772420100)
        candidate = {
            "window": [neighbour, anchor],
            "anchor_local_ids": [2],
        }
        chosen = build_table.best_window_message(candidate)
        self.assertEqual(chosen["local_id"], 2)

    def test_falls_back_to_whole_window(self):
        candidate = {"window": [make_msg(1, "通知：明天开会")], "anchor_local_ids": [99]}
        self.assertEqual(build_table.best_window_message(candidate)["local_id"], 1)

    def test_empty_window(self):
        self.assertEqual(build_table.best_window_message({"window": []}), {})


class TestBuildRecords(unittest.TestCase):
    def _candidate(self):
        return {
            "candidate_id": "c0001",
            "contact_display": "测试群",
            "chat_type": "group",
            "anchor_time": "2026-03-02 09:28:00",
            "anchor_local_ids": [1],
            "score": 9,
            "hits": {"keywords": ["通知"], "strong": ["通知"], "categories": ["通知类"],
                     "time": ["时点"], "actions": ["祈使"], "markers": []},
            "window": [make_msg(1, "通知一下：本周五下午3点在B301开班会，请全体成员务必参加。")],
        }

    def test_rule_mode(self):
        records = build_table.build_records([self._candidate()], None)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["会话类型"], "群聊")
        self.assertEqual(records[0]["通知类型"], "通知类")
        self.assertTrue(records[0]["置信度"].startswith("规则"))

    def test_ai_mode_overrides(self):
        notice = {
            "source_candidate": "c0001",
            "notice_time": "2026-03-02 09:28:00",
            "notice_type": "会议",
            "title": "周五班会",
            "event_time": "2026-03-06 15:00",
            "deadline": "",
            "location": "B301",
            "audience": "全体成员",
            "action_required": "需参加",
            "confidence": "high",
            "evidence": "原文",
        }
        records = build_table.build_records([self._candidate()], [notice])
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["通知类型"], "会议")
        self.assertEqual(records[0]["事项"], "周五班会")
        self.assertEqual(records[0]["截止时间"], "")
        self.assertEqual(records[0]["置信度"], "high")

    def test_ai_mode_ignores_unknown_candidate(self):
        records = build_table.build_records([self._candidate()],
                                            [{"source_candidate": "不存在", "title": "x"}])
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["事项"], "x")

    def test_headers_match_columns(self):
        record = build_table.build_records([self._candidate()], None)[0]
        for name, _, _ in build_table.COLUMNS:
            self.assertIn(name, record)


class TestXlsxWriter(unittest.TestCase):
    def test_writes_readable_workbook(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl 未安装")
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.xlsx")
            xlsx_min.write_xlsx(
                path, "通知表", ["序号", "事项", "原文"],
                [[1, "开会", "通知一下：明天下午3点开会"],
                 [2, "报名", "接龙报名春游"]],
                col_widths=[6, 20, 40], wrap_columns={1, 2},
            )
            self.assertTrue(os.path.exists(path))
            workbook = load_workbook(path)
            sheet = workbook["通知表"]
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A1"].value, "序号")
            self.assertEqual(sheet["B3"].value, "报名")
            self.assertEqual(sheet.freeze_panes, "A2")

    def test_escapes_illegal_xml_chars(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl 未安装")
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.xlsx")
            xlsx_min.write_xlsx(path, "s", ["a"], [["bad\x07char & <tag>"]])
            workbook = load_workbook(path)
            self.assertEqual(workbook["s"]["A2"].value, "badchar & <tag>")

    def test_sheet_name_truncated(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.xlsx")
            xlsx_min.write_xlsx(path, "x" * 60, ["a"], [["1"]])
            self.assertTrue(os.path.exists(path))


class TestEndToEnd(unittest.TestCase):
    def test_full_pipeline(self):
        with tempfile.TemporaryDirectory() as tmp:
            contacts = os.path.join(tmp, "data", "contacts")
            os.makedirs(contacts)
            demo_data.write_bundle(contacts, "群__t1", demo_data.build_payload(
                "测试群", "g1", "group", demo_data.GROUP_MESSAGES))
            demo_data.write_bundle(contacts, "张三__t2", demo_data.build_payload(
                "张三", "p1", "private", demo_data.PRIVATE_MESSAGES))

            bundles = scan_notices.discover_bundles(contacts)
            self.assertEqual(len(bundles), 2)

            all_candidates = []
            for name, path in bundles:
                payload = scan_notices.load_json(path)
                result = build_candidates(payload)
                all_candidates.extend(result["candidates"])
            for index, candidate in enumerate(all_candidates, start=1):
                candidate["candidate_id"] = "c%04d" % index

            self.assertGreaterEqual(len(all_candidates), 6)

            records = build_table.build_records(all_candidates, None)
            self.assertEqual(len(records), len(all_candidates))
            for record in records:
                self.assertTrue(record["通知时间"])
                self.assertTrue(record["会话"])

            out = os.path.join(tmp, "out.xlsx")
            headers = [n for n, _, _ in build_table.COLUMNS]
            widths = [w for _, w, _ in build_table.COLUMNS]
            flags = [f for _, _, f in build_table.COLUMNS]
            rows = [[r[n] for n in headers] for r in records]
            build_table.write_with_fallback(out, headers, rows, widths, flags, [])
            self.assertTrue(os.path.exists(out))
            self.assertGreater(os.path.getsize(out), 1000)

    def test_no_false_positive_on_chatter(self):
        """闲聊不应产生候选。"""
        chatter = [
            ("them", "张三", "text", "在吗"), ("me", "我", "text", "在"),
            ("them", "张三", "text", "哈哈昨天那个视频太好笑了"), ("me", "我", "text", "确实"),
            ("them", "张三", "text", "晚上一起吃饭吗"), ("me", "我", "text", "不了"),
            ("them", "张三", "text", "同学们早上好"), ("me", "我", "text", "早"),
        ]
        payload = demo_data.build_payload("闲聊", "x", "private", chatter)
        result = build_candidates(payload)
        self.assertEqual(len(result["candidates"]), 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
