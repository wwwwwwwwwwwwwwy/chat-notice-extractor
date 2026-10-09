"""build_table.py - 把候选/AI 精判结果渲染成 Excel 表格

两种输入模式：
  1. 仅规则候选（--candidates）        → 直接出一张「规则初判表」，用于快速预览
  2. 规则候选 + AI 精判（--notices）   → 出最终表，未通过精判的候选不进入主表

输出：
  out/notices.xlsx   主表（含「通知表」「候选明细」两个工作表）
  out/notices.csv    可选（--csv），UTF-8 with BOM，便于 Excel 直接打开
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from notice_rules import score_message  # noqa: E402

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# 主表列定义：(表头, 宽度, 是否自动换行)
COLUMNS = [
    ("序号", 6, False),
    ("通知时间", 19, False),
    ("会话", 16, False),
    ("会话类型", 9, False),
    ("发送人", 12, False),
    ("通知类型", 10, False),
    ("事项", 26, True),
    ("发生时间", 18, False),
    ("截止时间", 18, False),
    ("地点", 14, False),
    ("面向对象", 12, False),
    ("需回应", 10, False),
    ("置信度", 8, False),
    ("原文", 60, True),
]

CHAT_TYPE_LABEL = {"group": "群聊", "private": "单聊", "unknown": "未知"}


def load_json(path):
    with open(path, encoding="utf-8-sig") as handle:
        return json.load(handle)


def clip(text, limit):
    text = str(text or "").replace("\r\n", "\n").strip()
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def best_window_message(candidate):
    """定位候选的「通知主体」消息。

    窗口会向前后各扩展 context 条作为上下文，若直接取全局最高分，可能误选到
    邻近的另一条通知。因此优先在 anchor_local_ids 指定的锚点里选最高分。
    """
    window = candidate.get("window") or []
    if not window:
        return {}
    anchor_ids = set(candidate.get("anchor_local_ids") or [])
    pool = [item for item in window if item.get("local_id") in anchor_ids] or window
    return max(pool, key=lambda item: (score_message(item)["score"], -len(str(item.get("content")))))


def rule_time_fragment(candidate):
    """从命中项里回填一个粗粒度的时间片段，供规则模式预览。"""
    labels = (candidate.get("hits") or {}).get("time") or []
    return "、".join(labels)


def build_records(candidates, notices):
    """把候选与 AI 精判结果合并成统一记录列表。"""
    by_id = {c["candidate_id"]: c for c in candidates}
    records = []

    if notices is not None:
        for index, notice in enumerate(notices, start=1):
            candidate = by_id.get(notice.get("source_candidate"), {})
            window_msg = best_window_message(candidate) if candidate else {}
            records.append({
                "序号": index,
                "通知时间": notice.get("notice_time") or candidate.get("anchor_time", ""),
                "会话": notice.get("contact_display") or candidate.get("contact_display", ""),
                "会话类型": CHAT_TYPE_LABEL.get(
                    notice.get("chat_type") or candidate.get("chat_type", "unknown"), "未知"),
                "发送人": notice.get("sender") or window_msg.get("sender_name", ""),
                "通知类型": notice.get("notice_type", ""),
                "事项": notice.get("title", ""),
                "发生时间": notice.get("event_time", ""),
                "截止时间": notice.get("deadline", ""),
                "地点": notice.get("location", ""),
                "面向对象": notice.get("audience", ""),
                "需回应": notice.get("action_required", ""),
                "置信度": notice.get("confidence", ""),
                "原文": clip(notice.get("evidence") or window_msg.get("content", ""), 800),
                "_candidate": notice.get("source_candidate", ""),
            })
        return records

    # 规则初判模式
    for index, candidate in enumerate(candidates, start=1):
        window_msg = best_window_message(candidate)
        hits = candidate.get("hits") or {}
        categories = hits.get("categories") or []
        records.append({
            "序号": index,
            "通知时间": candidate.get("anchor_time", ""),
            "会话": candidate.get("contact_display", ""),
            "会话类型": CHAT_TYPE_LABEL.get(candidate.get("chat_type", "unknown"), "未知"),
            "发送人": window_msg.get("sender_name", ""),
            "通知类型": categories[0] if categories else "待判定",
            "事项": clip(window_msg.get("content", ""), 40),
            "发生时间": rule_time_fragment(candidate),
            "截止时间": "",
            "地点": "",
            "面向对象": "",
            "需回应": "",
            "置信度": f"规则{candidate.get('score', 0)}分",
            "原文": clip(window_msg.get("content", ""), 800),
            "_candidate": candidate.get("candidate_id", ""),
        })
    return records


def write_with_openpyxl(path, headers, rows, widths, wrap_flags, detail_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="2F5597")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")

    def fill_sheet(sheet, header_row, body):
        sheet.append(header_row)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in body:
            sheet.append(row)
        for index, (width, do_wrap) in enumerate(zip(widths, wrap_flags), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell, do_wrap in zip(row, wrap_flags):
                cell.alignment = wrap if do_wrap else top
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(header_row))}{max(sheet.max_row, 1)}"

    main = workbook.active
    main.title = "通知表"
    fill_sheet(main, headers, rows)

    if detail_rows:
        detail = workbook.create_sheet("候选明细")
        detail_headers = ["候选ID", "会话", "会话类型", "锚点时间", "规则分", "命中类别", "命中关键词", "窗口条数"]
        detail_widths = [10, 16, 9, 19, 8, 20, 34, 9]
        fill_sheet(detail, detail_headers, detail_rows)
        for index, width in enumerate(detail_widths, start=1):
            detail.column_dimensions[get_column_letter(index)].width = width

    workbook.save(path)


def write_with_fallback(path, headers, rows, widths, wrap_flags, detail_rows):
    from xlsx_min import write_xlsx

    write_xlsx(path, "通知表", headers, rows, col_widths=widths,
               wrap_columns={i for i, flag in enumerate(wrap_flags) if flag})

    if detail_rows:
        detail_path = path.replace(".xlsx", "__候选明细.xlsx")
        detail_headers = ["候选ID", "会话", "会话类型", "锚点时间", "规则分", "命中类别", "命中关键词", "窗口条数"]
        write_xlsx(detail_path, "候选明细", detail_headers, detail_rows,
                   col_widths=[10, 16, 9, 19, 8, 20, 34, 9])
        return detail_path
    return None


def main():
    parser = argparse.ArgumentParser(description="把通知候选渲染成 Excel 表格")
    parser.add_argument("--candidates", required=True, help="scan_notices.py 产出的 candidates.json")
    parser.add_argument("--notices", help="AI 精判产出的 notices.json（可选）")
    parser.add_argument("--out", required=True, help="输出 .xlsx 路径")
    parser.add_argument("--csv", action="store_true", help="同时输出 UTF-8-BOM 的 csv")
    args = parser.parse_args()

    candidates_doc = load_json(args.candidates)
    candidates = candidates_doc.get("candidates", [])

    notices = None
    if args.notices:
        if not os.path.isfile(args.notices):
            print(json.dumps({"status": "error",
                              "error": f"notices.json 不存在: {args.notices}"}, ensure_ascii=False),
                  file=sys.stderr)
            sys.exit(1)
        raw = load_json(args.notices)
        notices = raw.get("notices", []) if isinstance(raw, dict) else raw
        if not isinstance(notices, list):
            print(json.dumps({"status": "error",
                              "error": "notices.json 顶层应为数组或含 notices 数组"},
                             ensure_ascii=False), file=sys.stderr)
            sys.exit(1)

    records = build_records(candidates, notices)

    headers = [name for name, _, _ in COLUMNS]
    widths = [width for _, width, _ in COLUMNS]
    wrap_flags = [flag for _, _, flag in COLUMNS]
    rows = [[record[name] for name in headers] for record in records]

    detail_rows = []
    for candidate in candidates:
        hits = candidate.get("hits") or {}
        detail_rows.append([
            candidate.get("candidate_id", ""),
            candidate.get("contact_display", ""),
            CHAT_TYPE_LABEL.get(candidate.get("chat_type", "unknown"), "未知"),
            candidate.get("anchor_time", ""),
            candidate.get("score", 0),
            "、".join(hits.get("categories") or []),
            "、".join(hits.get("keywords") or [])[:200],
            len(candidate.get("window") or []),
        ])

    out_dir = os.path.dirname(os.path.abspath(args.out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    engine = "openpyxl"
    try:
        write_with_openpyxl(args.out, headers, rows, widths, wrap_flags, detail_rows)
    except ImportError:
        engine = "stdlib"
        write_with_fallback(args.out, headers, rows, widths, wrap_flags, detail_rows)

    csv_path = ""
    if args.csv:
        csv_path = os.path.splitext(args.out)[0] + ".csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)

    print(json.dumps({
        "status": "ok",
        "engine": engine,
        "mode": "ai+rule" if notices is not None else "rule-only",
        "rows": len(rows),
        "candidates": len(candidates),
        "output": os.path.abspath(args.out),
        "csv": os.path.abspath(csv_path) if csv_path else "",
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
